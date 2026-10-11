import csv
import io
import os
import re
import time
from typing import Any
import docx
from langchain_core.documents import Document
from langchain_core.vectorstores import InMemoryVectorStore
from langchain_text_splitters import RecursiveCharacterTextSplitter
import pypdf
from knowledge_base import get_embeddings

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".csv"}
MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024
MAX_FILES_PER_SESSION = 10
MAX_PAGES_PER_PDF = 100
MAX_CHUNKS_PER_DOC = 250
MAX_CSV_ROWS = 2500

BLOCKED_MAGIC_SIGNATURES = [
    b"MZ",
    b"\x7fELF",
    b"\xca\xfe\xba\xbe",
    b"\xfe\xed\xfa\xce",
    b"\xfe\xed\xfa\xcf",
    b"#!",
]


def sanitize_filename(filename: str) -> str:
    base_name = os.path.basename(filename).split("\\")[-1]
    name, ext = os.path.splitext(base_name)
    ext = ext.lower().strip()
    safe_name = re.sub(r"[^a-zA-Z0-9_\-]", "_", name).strip("._")
    if not safe_name:
        safe_name = "uploaded_document"
    return f"{safe_name[:64]}{ext}"


def validate_file_security(filename: str, file_bytes: bytes) -> tuple[bool, str, str]:
    if not file_bytes or len(file_bytes) == 0:
        return False, "", "File is empty (0 bytes). Please upload a valid document."

    if len(file_bytes) > MAX_FILE_SIZE_BYTES:
        max_mb = MAX_FILE_SIZE_BYTES / (1024 * 1024)
        return (
            False,
            "",
            f"File exceeds maximum allowed size ({len(file_bytes) / (1024*1024):.2f} MB > {max_mb:.0f} MB).",
        )

    safe_name = sanitize_filename(filename)
    _, ext = os.path.splitext(safe_name)

    if ext not in ALLOWED_EXTENSIONS:
        allowed_str = ", ".join(sorted(ALLOWED_EXTENSIONS))
        return (
            False,
            "",
            f"File extension '{ext}' is not permitted. Supported formats: {allowed_str}.",
        )

    for sig in BLOCKED_MAGIC_SIGNATURES:
        if file_bytes.startswith(sig):
            return (
                False,
                "",
                "Security Alert: File contains executable/script signatures and cannot be accepted.",
            )

    if ext == ".pdf":
        if not file_bytes.startswith(b"%PDF-"):
            return (
                False,
                "",
                "Invalid PDF format: File header does not match a genuine PDF document.",
            )

    elif ext == ".docx":
        if not file_bytes.startswith(b"PK\x03\x04"):
            return (
                False,
                "",
                "Invalid Word DOCX format: File header does not match a genuine Office document.",
            )

    elif ext in {".txt", ".md", ".csv"}:
        if b"\x00" in file_bytes[:4096]:
            return (
                False,
                "",
                "Security Alert: File appears to contain binary data instead of valid text.",
            )

    return True, safe_name, ""


def extract_text_from_file(file_bytes: bytes, safe_name: str) -> list[Document]:
    _, ext = os.path.splitext(safe_name)
    docs: list[Document] = []

    if ext == ".pdf":
        try:
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            if reader.is_encrypted:
                raise ValueError("Password-protected PDFs cannot be processed for security reasons.")

            total_pages = len(reader.pages)
            pages_to_process = min(total_pages, MAX_PAGES_PER_PDF)

            for page_idx in range(pages_to_process):
                page = reader.pages[page_idx]
                text = page.extract_text() or ""
                clean_text = text.strip()
                if clean_text:
                    docs.append(
                        Document(
                            page_content=clean_text,
                            metadata={
                                "source": safe_name,
                                "page": page_idx + 1,
                                "total_pages": total_pages,
                                "type": "pdf",
                            },
                        )
                    )
        except Exception as e:
            raise ValueError(f"Failed to safely parse PDF: {e}") from e

    elif ext == ".docx":
        try:
            doc = docx.Document(io.BytesIO(file_bytes))
            full_text_parts: list[str] = []

            for p in doc.paragraphs:
                p_text = p.text.strip()
                if p_text:
                    full_text_parts.append(p_text)

            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        full_text_parts.append(" | ".join(row_text))

            combined = "\n\n".join(full_text_parts)
            if combined.strip():
                docs.append(
                    Document(
                        page_content=combined,
                        metadata={
                            "source": safe_name,
                            "type": "docx",
                        },
                    )
                )
        except Exception as e:
            raise ValueError(f"Failed to safely parse Word DOCX document: {e}") from e

    elif ext in {".txt", ".md"}:
        try:
            text = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            try:
                text = file_bytes.decode("latin-1")
            except Exception as e:
                raise ValueError(f"Failed to decode text file: {e}") from e

        if text.strip():
            docs.append(
                Document(
                    page_content=text.strip(),
                    metadata={
                        "source": safe_name,
                        "type": ext[1:],
                    },
                )
            )

    elif ext == ".csv":
        try:
            try:
                decoded = file_bytes.decode("utf-8")
            except UnicodeDecodeError:
                decoded = file_bytes.decode("latin-1")

            reader = csv.reader(io.StringIO(decoded))
            rows = []
            header = None
            row_count = 0

            for row in reader:
                row_count += 1
                if row_count > MAX_CSV_ROWS:
                    break
                if not row or not any(field.strip() for field in row):
                    continue
                if header is None:
                    header = row
                    rows.append(f"Columns: {', '.join(header)}")
                else:
                    row_repr = ", ".join(
                        f"{header[i] if i < len(header) else f'Col{i+1}'}: {val.strip()}"
                        for i, val in enumerate(row)
                    )
                    rows.append(row_repr)

            csv_text = "\n".join(rows)
            if csv_text.strip():
                docs.append(
                    Document(
                        page_content=csv_text.strip(),
                        metadata={
                            "source": safe_name,
                            "type": "csv",
                            "row_count": min(row_count, MAX_CSV_ROWS),
                        },
                    )
                )
        except Exception as e:
            raise ValueError(f"Failed to safely parse CSV document: {e}") from e

    return docs


def build_user_vector_store(
    documents: list[Document],
    existing_store: InMemoryVectorStore | None = None,
    chunk_size: int = 1500,
    chunk_overlap: int = 150,
) -> tuple[InMemoryVectorStore, int]:
    if not documents:
        embeddings = get_embeddings()
        return existing_store or InMemoryVectorStore(embeddings), 0

    text_splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size, chunk_overlap=chunk_overlap
    )
    all_splits = text_splitter.split_documents(documents)

    if len(all_splits) > MAX_CHUNKS_PER_DOC:
        all_splits = all_splits[:MAX_CHUNKS_PER_DOC]

    for idx, split in enumerate(all_splits):
        split.metadata["source_type"] = "user_upload"
        split.metadata["chunk_id"] = idx + 1

    embeddings = get_embeddings()
    vector_store = existing_store if existing_store is not None else InMemoryVectorStore(embeddings)

    batch_size = 25
    total_splits = len(all_splits)

    for i in range(0, total_splits, batch_size):
        batch = all_splits[i : i + batch_size]
        max_retries = 4
        success = False

        for attempt in range(max_retries):
            try:
                vector_store.add_documents(documents=batch)
                success = True
                time.sleep(1.0)
                break
            except Exception as e:
                err_msg = str(e)
                if any(x in err_msg for x in ["RESOURCE_EXHAUSTED", "429", "quota", "503"]):
                    wait_time = 15 * (attempt + 1)
                    time.sleep(wait_time)
                else:
                    time.sleep(3)

        if not success:
            raise RuntimeError(
                f"Embedding API temporarily unavailable or quota exceeded while indexing batch {i} to {i + batch_size}."
            )

    return vector_store, total_splits
